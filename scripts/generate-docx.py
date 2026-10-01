#!/usr/bin/env python3
"""Génère un rapport Word à partir du markdown rapport-projet-complet.md."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "rapport" / "rapport-projet-complet.md"
OUT = ROOT / "docs" / "livrables" / "Rapport-Learning-Analytics.docx"


def add_heading(doc: Document, text: str, level: int) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x0F, 0x2C, 0x4C)


def main() -> None:
    text = SRC.read_text(encoding="utf-8")
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("Learning Analytics of Educational Processes", 0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x0F, 0x2C, 0x4C)

    p = doc.add_paragraph(
        "Équipe 1 — Cloud Computing I — SIGLIS M1 UPPA\n"
        "Bastien Bruey, Edouard Clemenceau, Titoan Lalanne, Jiale Li, Yoan Minbielle"
    )
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("# "):
            # already have main title
            continue
        if line.startswith("## "):
            add_heading(doc, line[3:].strip(), 1)
            continue
        if line.startswith("### "):
            add_heading(doc, line[4:].strip(), 2)
            continue
        if line.startswith("---"):
            continue
        if line.startswith("|") and "---" in line:
            continue
        if line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            doc.add_paragraph(" · ".join(cells), style="List Bullet")
            continue
        if line.startswith("```"):
            continue
        if line.startswith("- "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
            continue
        if line.startswith("> "):
            qp = doc.add_paragraph(line[2:].strip())
            qp.runs[0].italic = True
            continue
        if line.startswith("**") and line.endswith("**") and line.count("**") == 2:
            add_heading(doc, line.strip("*"), 2)
            continue
        # strip simple markdown bold markers lightly
        clean = line.replace("**", "")
        doc.add_paragraph(clean)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"DOCX: {OUT}")


if __name__ == "__main__":
    main()
