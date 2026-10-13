#!/usr/bin/env python3
"""PPTX — Présentation d'avancement 13/14 oct (PoC terminée)."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

NAVY = RGBColor(0x0F, 0x2C, 0x4C)
TEAL = RGBColor(0x1F, 0x6F, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF4, 0xF7, 0xFA)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x4A, 0x55, 0x68)

SLIDES = [
    (
        "Learning Analytics of Educational Processes",
        [
            "Équipe 1 — Cloud Computing I — SIGLIS M1",
            "Bastien Bruey · Edouard Clemenceau · Titoan Lalanne · Jiale Li · Yoan Minbielle",
            "Présentation d’avancement — PoC Docker TERMINÉE",
        ],
    ),
    (
        "Message d’avancement",
        [
            "Minimum demandé : problème + idée + début Docker",
            "Notre avancement : la PoC complète fonctionne déjà",
            "Aujourd’hui : on montre le résultat technique",
            "Suite : polish rapport / soutenance finale",
        ],
    ),
    (
        "Problématique",
        [
            "Traces d’apprentissage dispersées (LMS, quiz, vidéos…)",
            "Décrochage détecté trop tard",
            "Enseignants peu outillés pour agir avec la donnée",
        ],
    ),
    (
        "Solution proposée",
        [
            "Service SaaS de Learning Analytics",
            "Collecte standard xAPI → entrepôt → analytics / ML",
            "Dashboard : risque + remédiations pour l’enseignant",
            "Déploiement private cloud EVA (UPPA)",
        ],
    ),
    (
        "Positionnement cloud (NIST)",
        [
            "Service model : SaaS",
            "Deployment : Private cloud EVA",
            "VM : m1-siglis-cc-03 — 10.3.16.179",
            "Consumer = enseignant · Provider = équipe / UPPA",
        ],
    ),
    (
        "Architecture Docker",
        [
            "xapi-generator → lrs-api → postgres",
            "analytics-worker (engagement + risque ML)",
            "analytics-api → dashboard enseignants",
            "Optionnel : Grafana / Portainer",
        ],
    ),
    (
        "Réutilisé vs développé",
        [
            "Réutilisé : PostgreSQL, nginx, Grafana, Portainer",
            "Développé : LRS xAPI, générateur, worker ML, API, dashboard",
            "Conforme à la consigne : ni from scratch, ni collage d’images",
        ],
    ),
    (
        "Ce qui tourne aujourd’hui",
        [
            "Stack Compose multi-conteneurs opérationnelle",
            "~2000–3000+ statements xAPI · 25 apprenants simulés",
            "Dashboard : KPIs, risque, remédiations",
            "Smoke test automatisé : PASSED",
            "Hub Docker : images poussées (edouardclemenceau/...)",
        ],
    ),
    (
        "Démo",
        [
            "Dashboard : http://localhost:8300",
            "API : http://localhost:8210/api/overview",
            "EVA (VPN) : http://10.3.16.179:8300",
            "→ Montrer les apprenants à risque + remédiations",
        ],
    ),
    (
        "Ops & contraintes EVA",
        [
            "1 vCPU / 2 Go RAM → stack allégée",
            "Proxy UPPA + NO_PROXY inter-services",
            "restart: on-failure:3 (pas de boucle illimitée)",
            "OpenVPN obligatoire pour accéder à la VM",
        ],
    ),
    (
        "Roadmap",
        [
            "✓ Problème, solution, Docker PoC",
            "✓ Démo locale + EVA documentée",
            "→ Soutenance finale : BMC, ArchiMate, cookbook, slides finales",
            "Perspectives : LMS réel, RGPD, Swarm/K8s",
        ],
    ),
    (
        "Merci — questions ?",
        [
            "GitHub : github.com/galaxie44/learning-analytics-educational-processes",
            "Équipe 1 · Learning Analytics · Cloud Computing I",
        ],
    ),
]


def add_bg(slide, color, prs):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    sp_tree = slide.shapes._spTree
    sp = shape._element
    sp_tree.remove(sp)
    sp_tree.insert(2, sp)


def style_title(tf, text, size=36, color=WHITE):
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = True
    p.font.color.rgb = color
    p.font.name = "Calibri"


def add_bullets(tf, lines, size=22, color=DARK):
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.name = "Calibri"
        p.space_after = Pt(10)


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    for idx, (title, bullets) in enumerate(SLIDES):
        slide = prs.slides.add_slide(blank)
        if idx == 0:
            add_bg(slide, NAVY, prs)
            bar = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(0), Inches(5.9), prs.slide_width, Inches(1.6)
            )
            bar.fill.solid()
            bar.fill.fore_color.rgb = TEAL
            bar.line.fill.background()
            box = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.5), Inches(2.5))
            style_title(box.text_frame, title, size=32, color=WHITE)
            sub = slide.shapes.add_textbox(Inches(0.8), Inches(6.1), Inches(11.5), Inches(1.2))
            add_bullets(sub.text_frame, bullets, size=15, color=WHITE)
        else:
            add_bg(slide, LIGHT, prs)
            accent = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), prs.slide_height
            )
            accent.fill.solid()
            accent.fill.fore_color.rgb = TEAL
            accent.line.fill.background()
            header = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.35)
            )
            header.fill.solid()
            header.fill.fore_color.rgb = NAVY
            header.line.fill.background()
            tbox = slide.shapes.add_textbox(Inches(0.7), Inches(0.35), Inches(12), Inches(0.8))
            style_title(tbox.text_frame, title, size=28, color=WHITE)
            body = slide.shapes.add_textbox(Inches(0.9), Inches(1.85), Inches(11.5), Inches(4.8))
            add_bullets(body.text_frame, bullets, size=22, color=DARK)
            footer = slide.shapes.add_textbox(Inches(0.7), Inches(6.9), Inches(12), Inches(0.4))
            tf = footer.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = f"Équipe 1 · Avancement · PoC terminée  |  {idx}/{len(SLIDES) - 1}"
            p.font.size = Pt(12)
            p.font.color.rgb = MUTED
            p.font.name = "Calibri"

    out = (
        Path(__file__).resolve().parents[1]
        / "docs"
        / "livrables"
        / "Presentation-Avancement-COMPLETE.pptx"
    )
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    # also overwrite main presentation name
    prs.save(out.parent / "Presentation-Learning-Analytics.pptx")
    print(f"PPTX: {out}")


if __name__ == "__main__":
    main()
